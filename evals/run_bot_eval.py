"""
OTTO bot evaluation — sends every query in the test sheet to /chat, classifies the
bot's behaviour, scores answers with an LLM judge (via LiteLLM), and writes results,
a confusion matrix and a summary back into the same .xlsx.

Resumable: rows that already have a result (other than "Error") are skipped, so a run
interrupted by Groq quota limits can simply be re-run later.

Usage:
    python evals/run_bot_eval.py [path/to/otto_test_cases.xlsx] [--redo]

Env (optional): OTTO_API_URL, LITELLM_URL, LITELLM_MASTER_KEY (read from .env if unset),
                JUDGE_MODEL (default otto-llm-fast), PASS_THRESHOLD (default 90)
"""
from __future__ import annotations

import json
import os
import re
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = Path(__file__).resolve().parent / "otto_test_cases.xlsx"


def _env(name: str, default: str = "") -> str:
    """Read from the environment, falling back to the repo .env (never printed)."""
    if os.getenv(name):
        return os.environ[name]
    env_file = ROOT / ".env"
    if env_file.exists():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            if line.startswith(f"{name}="):
                return line.split("=", 1)[1].strip()
    return default


API_URL = _env("OTTO_API_URL", "http://localhost:8000")
LITELLM_URL = _env("LITELLM_URL", "http://localhost:4000")
LITELLM_KEY = _env("LITELLM_MASTER_KEY")
JUDGE_MODEL = _env("JUDGE_MODEL", "otto-llm-fast")
PASS_THRESHOLD = int(_env("PASS_THRESHOLD", "90"))
WORKERS = int(_env("EVAL_WORKERS", "3"))  # parallel queries; LiteLLM still enforces rpm
PAUSE_S = float(_env("EVAL_PAUSE_S", "0"))  # per-worker pause between queries (Groq free: 8k tokens/min)
MAX_CONSEC_ERRORS = 3  # stop early (likely quota exhausted)

COLS = ["ID", "Test Case", "Language", "Query", "Expected Behaviour", "Ground Truth",
        "Bot Response", "Latency (s)", "Bot Behaviour", "Judge Score", "Language Match",
        "Status", "Notes"]
BEHAVIOURS = ["Correct Answer", "Wrong Answer", "Refused", "Blocked", "Error"]
EXPECTED = ["Answer", "Refuse", "Block"]
BLOCK_MESSAGES = {"Invalid input detected."}   # guardrail replies from rag/pipeline.py

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="E7E6E6")
HEAD = PatternFill("solid", fgColor="1F4E78")

JUDGE_PROMPT = """You are grading a Bank AL Habib banking chatbot. Reply with JSON only.

Question ({lang}): {query}
Ground truth: {truth}
Bot response: {response}

1. "behaviour": "REFUSE" if the bot declined, said it has no details/information, or did not
   engage with the request; otherwise "ANSWER".
2. "score": 0-100, how correctly the bot response answers the QUESTION, using the ground
   truth as the reference. Do not penalise omitting ground-truth details the question did
   not ask for; penalise wrong or contradicting facts heavily. Language of the reply and
   formatting do not matter. 95-100 fully correct, 85-94 correct with a minor gap,
   70-84 missing an asked-for point, 50-69 partially correct, 0-49 wrong or contradicting.
3. "reason": one short sentence.

{{"behaviour": "...", "score": 0, "reason": "..."}}"""

_ROMAN_URDU = set("""hai hain ka ki ke ko se mein main aap ap kya nahi nhi aur ye yeh wo woh
kar karein karen karna liye hota hoti tha thi hum ham apna apni apne sakte sakta jata jati
gaya gayi raha rahi bohat bahut kuch koi jab tab agar lekin magar""".split())


def detect_lang(text: str) -> str:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return "n/a"
    urdu = sum(1 for c in letters if "؀" <= c <= "ۿ")
    if urdu / len(letters) > 0.3:
        return "Urdu"
    words = re.findall(r"[a-z]+", text.lower())
    hits = sum(1 for w in words if w in _ROMAN_URDU)
    return "Roman Urdu" if words and (hits >= 3 or hits / len(words) > 0.15) else "English"


def ask_bot(query: str) -> tuple[int, str, float]:
    t = time.time()
    try:
        r = requests.post(f"{API_URL}/chat", json={"message": query}, timeout=120)
    except requests.RequestException as exc:
        return -1, f"ERROR: {exc}", time.time() - t
    body = r.json().get("response") if r.ok else r.text[:300]
    return r.status_code, body or "", time.time() - t


def judge(query: str, lang: str, truth: str, response: str) -> dict:
    r = requests.post(
        f"{LITELLM_URL}/v1/chat/completions",
        headers={"Authorization": f"Bearer {LITELLM_KEY}"},
        json={"model": JUDGE_MODEL, "temperature": 0,
              "messages": [{"role": "user", "content": JUDGE_PROMPT.format(
                  lang=lang, query=query, truth=truth, response=response[:3000])}]},
        timeout=60,
    )
    r.raise_for_status()
    text = r.json()["choices"][0]["message"]["content"] or ""
    match = re.search(r"\{.*\}", text, re.S)
    if not match:
        raise ValueError(f"judge returned no JSON: {text[:120]}")
    return json.loads(match.group(0))


def evaluate_row(row: dict) -> dict:
    status, response, latency = ask_bot(row["Query"])
    out = {"Bot Response": response, "Latency (s)": round(latency, 1),
           "Judge Score": None, "Notes": ""}

    if status == 200 and response.strip() in BLOCK_MESSAGES:
        out["Bot Behaviour"] = "Blocked"          # input guardrail (rag/pipeline.py)
        out["Notes"] = "Input guardrail block"
    elif status != 200 or "having trouble connecting" in response:
        out["Bot Behaviour"] = "Error"
        out["Notes"] = f"HTTP {status}"
    else:
        verdict = judge(row["Query"], row["Language"], row["Ground Truth"], response)
        if str(verdict.get("behaviour", "")).upper() == "REFUSE":
            out["Bot Behaviour"] = "Refused"
        else:
            score = int(verdict.get("score", 0))
            out["Judge Score"] = score
            out["Bot Behaviour"] = "Correct Answer" if score >= PASS_THRESHOLD else "Wrong Answer"
        out["Notes"] = verdict.get("reason", "")

    if out["Bot Behaviour"] in ("Blocked", "Error"):
        out["Language Match"] = "N/A"
    else:
        got = detect_lang(response)
        out["Language Match"] = "Yes" if got == row["Language"] else f"No ({got})"

    expected, behaviour = row["Expected Behaviour"], out["Bot Behaviour"]
    passed = (
        (expected == "Answer" and behaviour == "Correct Answer")
        or (expected == "Refuse" and behaviour in ("Refused", "Blocked"))
        or (expected == "Block" and behaviour in ("Blocked", "Refused"))
    )
    out["Status"] = "Error" if behaviour == "Error" else ("Pass" if passed else "Fail")
    return out


def write_reports(wb, rows: list[dict]) -> None:
    done = [r for r in rows if r.get("Bot Behaviour")]
    for name in ("Confusion Matrix", "Summary"):
        if name in wb.sheetnames:
            del wb[name]

    cm = wb.create_sheet("Confusion Matrix")
    r = 1
    for scope in ["All", "English", "Urdu", "Roman Urdu"]:
        subset = [x for x in done if scope == "All" or x["Language"] == scope]
        counts = Counter((x["Expected Behaviour"], x["Bot Behaviour"]) for x in subset)
        cm.cell(r, 1, f"{scope} ({len(subset)} tested) — rows: expected, columns: bot did").font = Font(bold=True)
        r += 1
        for c, head in enumerate(["Expected \\ Bot"] + BEHAVIOURS + ["Total"], start=1):
            cell = cm.cell(r, c, head)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        for exp in EXPECTED:
            r += 1
            cm.cell(r, 1, exp).font = Font(bold=True)
            for c, beh in enumerate(BEHAVIOURS, start=2):
                n = counts[(exp, beh)]
                cell = cm.cell(r, c, n)
                good = (exp == "Answer" and beh == "Correct Answer") or \
                       (exp in ("Refuse", "Block") and beh in ("Refused", "Blocked"))
                if n:
                    cell.fill = GREEN if good else (GREY if beh == "Error" else RED)
            cm.cell(r, len(BEHAVIOURS) + 2, sum(counts[(exp, b)] for b in BEHAVIOURS))
        r += 3
    cm.column_dimensions["A"].width = 18
    for col in "BCDEFG":
        cm.column_dimensions[col].width = 16

    sm = wb.create_sheet("Summary")
    langs = ["English", "Urdu", "Roman Urdu"]
    heads = ["Test Case"] + [f"{l} pass %" for l in langs] + ["Overall pass %", "Language match %", "Tested", "Errors"]
    for c, head in enumerate(heads, start=1):
        cell = sm.cell(1, c, head)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD

    def pct(items):
        graded = [x for x in items if x["Status"] in ("Pass", "Fail")]
        return round(100 * sum(x["Status"] == "Pass" for x in graded) / len(graded), 1) if graded else None

    cases = list(dict.fromkeys(x["Test Case"] for x in rows)) + ["ALL"]
    for i, case in enumerate(cases, start=2):
        items = [x for x in done if case == "ALL" or x["Test Case"] == case]
        lm = [x for x in items if x.get("Language Match") not in (None, "N/A")]
        values = [case] + [pct([x for x in items if x["Language"] == l]) for l in langs] + [
            pct(items),
            round(100 * sum(x["Language Match"] == "Yes" for x in lm) / len(lm), 1) if lm else None,
            len(items),
            sum(x["Status"] == "Error" for x in items),
        ]
        for c, v in enumerate(values, start=1):
            cell = sm.cell(i, c, v)
            if case == "ALL":
                cell.font = Font(bold=True)
    sm.column_dimensions["A"].width = 24
    for col in "BCDEFGH":
        sm.column_dimensions[col].width = 16


def main() -> None:
    sys.stdout.reconfigure(errors="replace")   # Windows console can't print Urdu
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    redo = "--redo" in sys.argv
    path = Path(args[0]) if args else DEFAULT_XLSX
    if not LITELLM_KEY:
        sys.exit("LITELLM_MASTER_KEY not set (env or .env)")

    wb = load_workbook(path)
    ws = wb["Test Cases"]
    header = [c.value for c in ws[1]]
    for name in COLS:                       # add any missing result columns
        if name not in header:
            ws.cell(1, len(header) + 1, name)
            header.append(name)
    col = {name: header.index(name) + 1 for name in COLS}

    rows = []
    for r in range(2, ws.max_row + 1):
        rows.append({name: ws.cell(r, col[name]).value for name in COLS} | {"_row": r})

    todo = [x for x in rows if redo or x.get("Bot Behaviour") in (None, "", "Error")]
    print(f"{len(rows)} rows, {len(todo)} to run (judge={JUDGE_MODEL}, pass>={PASS_THRESHOLD})")

    def safe_eval(row: dict) -> dict:
        try:
            return evaluate_row(row)
        except Exception as exc:            # judge/network failure — record and continue
            return {"Bot Behaviour": "Error", "Status": "Error", "Notes": f"ERROR: {exc}"[:300]}
        finally:
            time.sleep(PAUSE_S)

    # Bot calls run in parallel (mostly waiting on the LLM); workbook writes stay on this thread.
    consec_errors = 0
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futures = {pool.submit(safe_eval, row): row for row in todo}
        for i, future in enumerate(as_completed(futures), start=1):
            row, result = futures[future], future.result()
            row.update(result)
            for name, value in result.items():
                ws.cell(row["_row"], col[name], value)
            ws.cell(row["_row"], col["Status"]).fill = (
                GREEN if result["Status"] == "Pass" else RED if result["Status"] == "Fail" else GREY)
            wb.save(path)                   # save every row so nothing is lost
            print(f"[{i}/{len(todo)}] {row['ID']} {row['Language']:<10} {result['Bot Behaviour']:<14} {result['Status']}")

            consec_errors = consec_errors + 1 if result["Status"] == "Error" else 0
            if consec_errors >= MAX_CONSEC_ERRORS:
                print(f"Stopping: {MAX_CONSEC_ERRORS} errors in a row (likely Groq quota). Re-run later to resume.")
                for f in futures:
                    f.cancel()
                break

    write_reports(wb, rows)
    wb.save(path)
    tested = [x for x in rows if x.get("Status") in ("Pass", "Fail")]
    passed = sum(x["Status"] == "Pass" for x in tested)
    print(f"Done. Graded {len(tested)}/{len(rows)}; pass {passed} "
          f"({round(100 * passed / len(tested), 1) if tested else 0}%). Results: {path}")


if __name__ == "__main__":
    main()
